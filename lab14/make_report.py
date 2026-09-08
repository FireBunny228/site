#!/usr/bin/env python3
"""Generate Lab 14 Word report."""

from docx import Document
from docx.shared import Pt, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from pathlib import Path

doc = Document()

# Page margins
for section in doc.sections:
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(1.5)

style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(14)


def add_center(text, bold=False, size=14, space_after=6):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    p.paragraph_format.space_after = Pt(space_after)
    return p


def add_para(text, bold=False, first_line=True, align="justify"):
    p = doc.add_paragraph()
    if align == "justify":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    elif align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "left":
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)
    if first_line:
        p.paragraph_format.first_line_indent = Cm(1.25)
    p.paragraph_format.line_spacing = 1.5
    return p


def add_heading_custom(text, level=1):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.first_line_indent = Cm(0)
    return p


# Title page-ish header
add_center("Лабораторная работа №14", bold=True, size=16)
add_center("Диаграмма деятельности (Activity Diagram)", bold=True, size=14)
add_center("Тема: Бронирование мест на экскурсии и туры:", bold=False, size=14)
add_center("даты, маршруты, количество мест", bold=False, size=14)
doc.add_paragraph()

add_heading_custom("1. Цель работы")
add_para(
    "Целью лабораторной работы является построение диаграммы деятельности UML "
    "для предметной области «Бронирование мест на экскурсии и туры». Диаграмма "
    "должна отражать последовательность действий участников процесса, точки принятия "
    "решений и параллельные действия при подтверждении бронирования."
)

add_heading_custom("2. Постановка задачи")
add_para(
    "Необходимо разработать диаграмму деятельности с разделением на дорожки "
    "(swimlanes), описывающую процесс бронирования мест на экскурсию или тур. "
    "Клиент выбирает дату, маршрут и количество мест; система проверяет доступность; "
    "при успешной регистрации заявки параллельно настраивается автонапоминание "
    "и уведомляется гид/туроператор; клиент получает подтверждение в личном кабинете."
)

add_heading_custom("3. Участники процесса (дорожки диаграммы)")

# Table of actors
table = doc.add_table(rows=4, cols=2)
table.style = "Table Grid"
table.alignment = WD_TABLE_ALIGNMENT.CENTER
headers = [("Участник", "Роль в процессе")]
data = [
    ("Клиент", "Авторизация, выбор даты, маршрута и количества мест, получение подтверждения"),
    ("Система/Админ", "Проверка доступности мест, регистрация заявки, настройка автонапоминания"),
    ("Гид/Туроператор", "Получение сообщения о новом бронировании"),
]
cells = table.rows[0].cells
cells[0].text = "Участник"
cells[1].text = "Роль в процессе"
for i, (a, b) in enumerate(data, start=1):
    table.rows[i].cells[0].text = a
    table.rows[i].cells[1].text = b
for row in table.rows:
    for cell in row.cells:
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(12)

doc.add_paragraph()

add_heading_custom("4. Описание процесса")
add_para(
    "Процесс начинается с авторизации клиента в системе. Далее клиент последовательно "
    "выбирает дату экскурсии или тура, маршрут и указывает необходимое количество мест. "
    "Эти параметры соответствуют предметной области лабораторной работы: даты, маршруты "
    "и количество мест."
)
add_para(
    "После выбора параметров система выполняет проверку доступности мест. "
    "В точке принятия решения «Свободно?» возможны два исхода:"
)
add_para("• Нет — клиент возвращается к выбору даты экскурсии/тура и повторяет подбор параметров;", first_line=False)
add_para("• Да — система регистрирует заявку на бронирование.", first_line=False)
add_para(
    "После регистрации заявки выполняется разветвление (fork) на два параллельных действия: "
    "система настраивает автоматическое напоминание, а гид/туроператор получает сообщение "
    "о бронировании. После синхронизации (join) клиент получает подтверждение в личном "
    "кабинете, и процесс завершается."
)

add_heading_custom("5. Диаграмма деятельности")
add_para("На рисунке 1 представлена диаграмма деятельности процесса бронирования.", first_line=False)

img_path = Path("/workspace/lab14/diagramma_deyatelnosti.png")
if img_path.exists():
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(img_path), width=Inches(5.8))

add_center("Рисунок 1 — Диаграмма деятельности: бронирование мест на экскурсии и туры", size=12)

add_heading_custom("6. Соответствие действий теме")
table2 = doc.add_table(rows=5, cols=2)
table2.style = "Table Grid"
table2.rows[0].cells[0].text = "Элемент темы"
table2.rows[0].cells[1].text = "Отражение на диаграмме"
rows2 = [
    ("Даты", "Действие «Выбрать дату экскурсии/тура»"),
    ("Маршруты", "Действие «Выбрать маршрут»"),
    ("Количество мест", "Действие «Указать количество мест» и проверка доступности"),
    ("Бронирование", "Регистрация заявки, уведомление гида, подтверждение клиенту"),
]
for i, (a, b) in enumerate(rows2, start=1):
    table2.rows[i].cells[0].text = a
    table2.rows[i].cells[1].text = b
for row in table2.rows:
    for cell in row.cells:
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(12)

doc.add_paragraph()

add_heading_custom("7. Вывод")
add_para(
    "В ходе лабораторной работы построена диаграмма деятельности UML для процесса "
    "бронирования мест на экскурсии и туры. Диаграмма отражает взаимодействие клиента, "
    "системы и гида/туроператора, проверку доступности мест, цикл повторного выбора "
    "при отсутствии свободных мест, а также параллельную обработку напоминания и "
    "уведомления. Полученная модель соответствует теме работы и может использоваться "
    "при проектировании информационной системы бронирования."
)

out = Path("/workspace/lab14/Lab14_Bronirovanie_ekskursiy.docx")
doc.save(out)
print("Saved", out)
